import os
import re
from io import BytesIO
from typing import Tuple


def parse_txt(content: str) -> Tuple[str, dict]:
    return content, {"type": "txt", "char_count": len(content)}


def parse_markdown(content: str) -> Tuple[str, dict]:
    text_content = re.sub(r'#+\s*', '', content)
    text_content = re.sub(r'\*\*([^*]+)\*\*', r'\1', text_content)
    text_content = re.sub(r'\*([^*]+)\*', r'\1', text_content)
    text_content = re.sub(r'`([^`]+)`', r'\1', text_content)
    text_content = re.sub(r'```[\s\S]*?```', '', text_content)
    text_content = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text_content)
    text_content = re.sub(r'\n{2,}', '\n', text_content)
    
    return text_content.strip(), {
        "type": "markdown",
        "char_count": len(text_content),
        "original_char_count": len(content)
    }


def parse_docx(content: bytes) -> Tuple[str, dict]:
    try:
        from docx import Document
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text_content = '\n\n'.join(paragraphs)
        
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    tables_text.append(' | '.join(row_text))
        
        if tables_text:
            text_content += '\n\n' + '\n'.join(tables_text)
        
        return text_content.strip(), {
            "type": "docx",
            "char_count": len(text_content),
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables)
        }
    except ImportError:
        return "", {"type": "docx", "char_count": 0, "error": "python-docx 未安装"}
    except Exception as e:
        return "", {"type": "docx", "char_count": 0, "error": str(e)}


def parse_document(file_content: str, filename: str) -> Tuple[str, dict]:
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.txt':
        return parse_txt(file_content)
    elif ext in ['.md', '.markdown']:
        return parse_markdown(file_content)
    else:
        return file_content, {"type": "unknown", "char_count": len(file_content)}


def parse_document_bytes(file_content: bytes, filename: str) -> Tuple[str, dict]:
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.docx':
        return parse_docx(file_content)
    elif ext == '.txt':
        try:
            return parse_txt(file_content.decode('utf-8'))
        except UnicodeDecodeError:
            return parse_txt(file_content.decode('gbk', errors='ignore'))
    elif ext in ['.md', '.markdown']:
        try:
            return parse_markdown(file_content.decode('utf-8'))
        except UnicodeDecodeError:
            return parse_markdown(file_content.decode('gbk', errors='ignore'))
    else:
        try:
            text_content = file_content.decode('utf-8')
        except UnicodeDecodeError:
            text_content = file_content.decode('gbk', errors='ignore')
        return text_content, {"type": "unknown", "char_count": len(text_content)}
