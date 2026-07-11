import pytest
from io import BytesIO
from app.documents.parser import (
    parse_txt,
    parse_markdown,
    parse_docx,
    parse_pdf,
    parse_document,
    parse_document_bytes
)


class TestParser:
    def test_parse_txt(self):
        content = "Hello World\nThis is a test."
        result, metadata = parse_txt(content)
        
        assert result == content
        assert metadata["type"] == "txt"
        assert metadata["char_count"] == len(content)

    def test_parse_markdown(self):
        content = "# Heading\n\n**Bold text** and *italic text*\n\n[Link](http://example.com)\n\n```code```"
        result, metadata = parse_markdown(content)
        
        assert "#" not in result
        assert "**" not in result
        assert "*" not in result
        assert "[" not in result
        assert "](" not in result
        assert "```" not in result
        assert metadata["type"] == "markdown"
        assert metadata["char_count"] > 0
        assert metadata["original_char_count"] == len(content)

    def test_parse_docx_valid(self):
        try:
            from docx import Document
            
            doc = Document()
            doc.add_paragraph("Hello World")
            doc.add_paragraph("This is a test document")
            
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            result, metadata = parse_docx(buffer.read())
            
            assert "Hello World" in result
            assert "test document" in result
            assert metadata["type"] == "docx"
            assert metadata["char_count"] > 0
            assert metadata["paragraph_count"] == 2
            assert metadata["table_count"] == 0
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_parse_docx_with_table(self):
        try:
            from docx import Document
            
            doc = Document()
            doc.add_paragraph("Document with table")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Test"
            table.cell(1, 1).text = "123"
            
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            result, metadata = parse_docx(buffer.read())
            
            assert "Name" in result
            assert "Value" in result
            assert "Test" in result
            assert "123" in result
            assert metadata["table_count"] == 1
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_parse_pdf_valid(self):
        try:
            import fitz
            
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Hello World")
            page.insert_text((72, 100), "This is a PDF test")
            
            buffer = BytesIO()
            buffer.write(doc.write())
            buffer.seek(0)
            
            result, metadata = parse_pdf(buffer.read())
            
            assert "Hello World" in result
            assert "PDF test" in result
            assert metadata["type"] == "pdf"
            assert metadata["char_count"] > 0
            assert metadata["page_count"] == 1
            assert metadata["is_text_based"] == True
        except ImportError:
            pytest.skip("PyMuPDF not installed")

    def test_parse_document_txt(self):
        content = "Test content"
        result, metadata = parse_document(content, "test.txt")
        
        assert result == content
        assert metadata["type"] == "txt"

    def test_parse_document_markdown(self):
        content = "# Title\n**bold**"
        result, metadata = parse_document(content, "test.md")
        
        assert metadata["type"] == "markdown"
        assert "#" not in result

    def test_parse_document_unknown(self):
        content = "Some content"
        result, metadata = parse_document(content, "test.unknown")
        
        assert result == content
        assert metadata["type"] == "unknown"

    def test_parse_document_bytes_txt(self):
        content = "Test content".encode('utf-8')
        result, metadata = parse_document_bytes(content, "test.txt")
        
        assert result == "Test content"
        assert metadata["type"] == "txt"

    def test_parse_document_bytes_docx(self):
        try:
            from docx import Document
            
            doc = Document()
            doc.add_paragraph("Test docx")
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            result, metadata = parse_document_bytes(buffer.read(), "test.docx")
            
            assert "Test docx" in result
            assert metadata["type"] == "docx"
        except ImportError:
            pytest.skip("python-docx not installed")

    def test_parse_document_bytes_pdf(self):
        try:
            import fitz
            
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Test PDF")
            
            buffer = BytesIO()
            buffer.write(doc.write())
            buffer.seek(0)
            
            result, metadata = parse_document_bytes(buffer.read(), "test.pdf")
            
            assert "Test PDF" in result
            assert metadata["type"] == "pdf"
        except ImportError:
            pytest.skip("PyMuPDF not installed")

    def test_parse_document_bytes_markdown(self):
        content = "# Title".encode('utf-8')
        result, metadata = parse_document_bytes(content, "test.md")
        
        assert metadata["type"] == "markdown"
        assert "#" not in result

    def test_parse_document_bytes_utf8(self):
        content = "测试内容".encode('utf-8')
        result, metadata = parse_document_bytes(content, "test.txt")
        
        assert result == "测试内容"
        assert metadata["type"] == "txt"
